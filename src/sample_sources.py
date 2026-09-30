#!/usr/bin/env python3
"""Create deterministic, disjoint raw training and held-out source samples."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


def record_id(source_id: str, row: dict[str, object], spec: dict) -> str:
    if spec.get("id_mode") == "row_fingerprint":
        payload = json.dumps(row, sort_keys=True, default=str, separators=(",", ":"))
        return hashlib.sha256(payload.encode()).hexdigest()
    values = [str(row.get(field, "")).strip() for field in spec["id_fields"]]
    if not all(values):
        raise ValueError(f"{source_id}: missing stable ID fields {spec['id_fields']}")
    return "|".join(values)


def first_text(element: ET.Element, tag: str) -> str:
    found = element.find(f".//{tag}")
    return (found.text or "").strip() if found is not None else ""


def xml_rows(path: Path, member: str) -> Iterator[dict[str, str]]:
    with zipfile.ZipFile(path) as archive, archive.open(member) as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            if element.tag == "ABR":
                yield {
                    "ABN": first_text(element, "ABN"), "ASICNumber": first_text(element, "ASICNumber"),
                    "EntityTypeText": first_text(element, "EntityTypeText"),
                    "NonIndividualNameText": first_text(element, "NonIndividualNameText"),
                    "State": first_text(element, "State"), "Postcode": first_text(element, "Postcode"),
                }
                element.clear()


def delimited_rows(handle: io.TextIOBase) -> Iterator[dict[str, str]]:
    sample = handle.read(8192)
    handle.seek(0)
    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    for row in csv.DictReader(handle, dialect=dialect):
        # DictReader uses None as a key for surplus fields in malformed rows.
        # Preserve such values under a deterministic, explicit key.
        yield {"_extra_fields" if key is None else str(key): value for key, value in row.items()}


def source_rows(source: dict, spec: dict, raw_dir: Path) -> Iterator[dict[str, object]]:
    path = raw_dir / source["local_filename"]
    reader = spec["reader"]
    if reader == "zip_xml":
        yield from xml_rows(path, spec["member"])
    elif reader == "zip_delimited":
        with zipfile.ZipFile(path) as archive, archive.open(spec["member"]) as binary:
            with io.TextIOWrapper(binary, encoding="utf-8-sig", newline="") as text:
                yield from delimited_rows(text)
    elif reader == "csv":
        with path.open(encoding="utf-8-sig", newline="") as text:
            yield from delimited_rows(text)
    elif reader == "xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[spec["sheet"]]
        rows = worksheet.iter_rows(values_only=True)
        for _ in range(spec.get("header_row", 1) - 1): next(rows)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        for values in rows:
            row = {headers[index]: value for index, value in enumerate(values) if headers[index]}
            if any(value is not None and value != "" for value in row.values()): yield row
    else:
        raise ValueError(f"Unsupported reader: {reader}")


def main() -> int:
    root = Path(".")
    sources = json.loads((root / "data/manifests/selected_sources.json").read_text())["sources"]
    specs = json.loads((root / "data/manifests/sampling_specs.json").read_text())
    sample_size = specs["sample_size_per_split"]
    raw_dir, sample_dir = root / "data/raw", root / "data/samples"
    manifest = {"sample_size_per_split": sample_size, "sources": []}
    for source in sources:
        source_id, spec = source["source_id"], specs["sources"][source["source_id"]]
        source["local_filename"] = next(path.name for path in raw_dir.iterdir() if path.stem == source_id)
        training, heldout, seen = [], [], set()
        for row in source_rows(source, spec, raw_dir):
            try: rid = record_id(source_id, row, spec)
            except ValueError: continue
            if rid in seen: continue
            seen.add(rid)
            target = training if int(hashlib.sha256(f"{source_id}:{rid}".encode()).hexdigest(), 16) % 2 == 0 else heldout
            if len(target) < sample_size: target.append({"source_record_id": rid, "raw": row})
            if len(training) == sample_size and len(heldout) == sample_size: break
        if len(training) < sample_size or len(heldout) < sample_size:
            raise RuntimeError(f"{source_id}: only training={len(training)}, heldout={len(heldout)} unique rows")
        destination = sample_dir / source_id
        destination.mkdir(parents=True, exist_ok=True)
        outputs = {}
        for name, rows in (("training", training), ("heldout", heldout)):
            path = destination / f"{name}.jsonl"
            path.write_text("".join(json.dumps(row, default=str) + "\n" for row in rows), encoding="utf-8")
            outputs[name] = {"path": str(path), "records": len(rows), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        manifest["sources"].append({
            "source_id": source_id,
            "stable_record_id": "sha256(canonical raw row)" if spec.get("id_mode") == "row_fingerprint" else spec["id_fields"],
            "outputs": outputs,
        })
        print(f"{source_id}: training={len(training)} heldout={len(heldout)}")
    output = root / "outputs/manifests/sampling_manifest.json"
    output.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    try: raise SystemExit(main())
    except (OSError, RuntimeError, ValueError, KeyError, StopIteration, csv.Error, zipfile.BadZipFile) as error:
        print(f"Sampling failed: {error}", file=sys.stderr); raise SystemExit(1)
